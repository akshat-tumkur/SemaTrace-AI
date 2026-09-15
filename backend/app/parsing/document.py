from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi import HTTPException
from fitz import open as open_pdf


@dataclass(frozen=True)
class ParsedDocument:
    filename: str
    text: str
    file_type: str


def parse_document(filename: str, content: bytes) -> ParsedDocument:
    extension = Path(filename).suffix.lower()
    if extension == ".txt":
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError as error:
            raise HTTPException(status_code=415, detail="Text files must be UTF-8 encoded.") from error
    elif extension == ".pdf":
        text = _parse_pdf(content)
    elif extension == ".docx":
        text = _parse_docx(content)
    else:
        raise HTTPException(status_code=415, detail="Supported files are .txt, .pdf, and .docx.")

    normalized = _normalize_text(text)
    if not normalized:
        raise HTTPException(status_code=422, detail="The uploaded document is empty.")
    return ParsedDocument(filename=filename, text=normalized, file_type=extension[1:])


def _parse_pdf(content: bytes) -> str:
    try:
        with open_pdf(stream=content, filetype="pdf") as document:
            return "\n".join(page.get_text() for page in document)
    except Exception as error:
        raise HTTPException(status_code=422, detail="The PDF could not be parsed.") from error


def _parse_docx(content: bytes) -> str:
    try:
        document = Document(BytesIO(content))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    except Exception as error:
        raise HTTPException(status_code=422, detail="The DOCX could not be parsed.") from error


def _normalize_text(text: str) -> str:
    lines = [" ".join(line.split()) for line in text.splitlines()]
    return "\n".join(line for line in lines if line).strip()
